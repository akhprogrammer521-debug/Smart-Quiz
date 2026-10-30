"""
Tests for the document summarization feature.

- Unit tests: fast, mocked (no API key needed).
- Integration tests: use real files in test_files/ and the real Gemini API.
  They print the actual summary so you can judge model quality.
"""

import os
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import httpx
import pytest
from google.genai import errors as genai_errors

import summarization
import summarization_api
from summarization import (
    DocConverterUnavailableError,
    ExtractedDocument,
    EmptyDocumentError,
    FileTooSmallError,
    GeminiConfigurationError,
    GeminiResponseError,
    GeminiUnavailableError,
    chunk_text,
    clean_text,
    extract_text_from_file,
    generate_final_summary,
    prepare_document_text,
    summarize_chunk,
    summarize_document,
    summarize_prepared_text,
)
from summarization_api import allowed_file, app

TEST_FILES_DIR = Path(__file__).parent / "test_files"
SAMPLE_TXT = TEST_FILES_DIR / "sample.txt"
SAMPLE_PDF = TEST_FILES_DIR / "sample.pdf"
SAMPLE_DOCX = TEST_FILES_DIR / "sample.docx"

# Skip real-API tests when the key is not available
HAS_API_KEY = bool(os.environ.get("GEMINI_API_KEY"))
skip_without_key = pytest.mark.skipif(
    not HAS_API_KEY,
    reason="GEMINI_API_KEY not set — skipping real Gemini integration tests",
)


# ---------------------------------------------------------------------------
# Unit tests (mocked, no API key required)
# ---------------------------------------------------------------------------


class TestCleanText:
    def test_collapses_extra_spaces(self):
        text = "Hello    world\t\tagain"
        assert clean_text(text) == "Hello world again"

    def test_collapses_blank_lines(self):
        text = "Line one\n\n\n\nLine two"
        assert clean_text(text) == "Line one\n\nLine two"

    def test_empty_string(self):
        assert clean_text("") == ""

    def test_does_not_rewrite_content(self):
        # Cleaning must not summarize or change wording
        text = "Photosynthesis converts light energy into chemical energy."
        assert "Photosynthesis" in clean_text(text)
        assert "chemical energy" in clean_text(text)


class TestChunkText:
    def test_empty_text(self):
        assert chunk_text("") == []

    def test_short_text_one_chunk(self):
        text = "Short paragraph about cells."
        chunks = chunk_text(text, chunk_size=2000)
        assert len(chunks) == 1
        assert chunks[0] == text

    def test_splits_by_paragraphs(self):
        # Two paragraphs that together exceed chunk_size
        p1 = "A" * 100
        p2 = "B" * 100
        text = f"{p1}\n\n{p2}"
        chunks = chunk_text(text, chunk_size=150)
        assert len(chunks) == 2
        assert chunks[0] == p1
        assert chunks[1] == p2

    def test_chunk_size_is_input_limit(self):
        paragraphs = [f"Paragraph {i} " + ("x" * 50) for i in range(10)]
        text = "\n\n".join(paragraphs)
        chunks = chunk_text(text, chunk_size=200)
        for chunk in chunks:
            # Allow a little slack only for the join characters inside a chunk
            assert len(chunk) <= 200 or "\n\n" in chunk


class TestExtractText:
    def test_extract_txt(self):
        text = extract_text_from_file(SAMPLE_TXT)
        assert "Photosynthesis" in text
        assert len(text) > 50

    def test_extract_pdf(self):
        text = extract_text_from_file(SAMPLE_PDF)
        assert "Photosynthesis" in text or "photosynthesis" in text.lower()

    def test_pdf_extraction_preserves_page_boundaries(self, tmp_path):
        import pymupdf

        source = tmp_path / "two-pages.pdf"
        pdf = pymupdf.open()
        pdf.new_page().insert_text((72, 72), "First page")
        pdf.new_page().insert_text((72, 72), "Second page")
        pdf.save(source)
        pdf.close()

        extracted = summarization.extract_document(source)
        assert extracted.page_boundaries[0] == 0
        assert len(extracted.page_boundaries) == 2
        assert extracted.text[extracted.page_boundaries[1]:].startswith("Second page")

    def test_extract_docx(self):
        text = extract_text_from_file(SAMPLE_DOCX)
        assert "Photosynthesis" in text or "photosynthesis" in text.lower()

    def test_docx_extraction_preserves_styled_headings(self, tmp_path):
        from docx import Document

        source = tmp_path / "headings.docx"
        document = Document()
        document.add_heading("Cell Biology", level=1)
        document.add_paragraph("Cells contain DNA.")
        document.save(source)

        extracted = summarization.extract_document(source)
        assert extracted.styled_headings == ("Cell Biology",)
        assert "Cells contain DNA." in extracted.text

    def test_doc_and_docx_extension_validation(self):
        assert allowed_file("lesson.doc")
        assert allowed_file("LESSON.DOCX")
        assert not allowed_file("lesson.doc.txt.exe")

    def test_docx_extraction_routing(self, tmp_path):
        source = tmp_path / "lesson.docx"
        source.write_bytes(b"test")
        with patch("summarization._extract_docx", return_value=ExtractedDocument("Word content")) as extract:
            assert summarization.extract_document(source).text == "Word content"
        extract.assert_called_once_with(source)

    def test_doc_conversion_routes_to_docx_extraction(self, tmp_path):
        source = tmp_path / "lesson.doc"
        source.write_bytes(b"legacy Word data")

        def fake_convert(args, **_kwargs):
            output_dir = Path(args[args.index("--outdir") + 1])
            (output_dir / "lesson.docx").write_bytes(b"converted Word data")
            return SimpleNamespace(returncode=0)

        with patch("summarization._find_libreoffice", return_value="soffice"), patch(
            "summarization.subprocess.run", side_effect=fake_convert
        ), patch("summarization._extract_docx", return_value=ExtractedDocument("Converted content")) as extract:
            assert summarization.extract_document(source).text == "Converted content"
            assert extract.call_args.args[0].suffix == ".docx"

    def test_doc_requires_real_converter(self, tmp_path):
        source = tmp_path / "lesson.doc"
        source.write_bytes(b"legacy Word data")
        with patch("summarization._find_libreoffice", side_effect=DocConverterUnavailableError("missing")):
            with pytest.raises(DocConverterUnavailableError):
                summarization.extract_document(source)

    def test_missing_file(self):
        with pytest.raises(FileNotFoundError):
            extract_text_from_file("does_not_exist.txt")

    def test_unsupported_type(self):
        fake = TEST_FILES_DIR / "sample.xyz"
        fake.write_text("hello", encoding="utf-8")
        try:
            with pytest.raises(ValueError):
                extract_text_from_file(fake)
        finally:
            fake.unlink(missing_ok=True)


class TestSummarizeWithMocks:
    @pytest.mark.parametrize("source", [SAMPLE_TXT, SAMPLE_PDF, SAMPLE_DOCX])
    @patch("summarization.os.path.getsize", return_value=10 * 1024)
    def test_prepare_supported_study_documents_without_gemini(self, _mock_size, source):
        result = prepare_document_text(source)
        assert result["original_text"].strip()
        assert result["normalized_text"].strip()

    @patch("summarization.extract_document")
    def test_file_smaller_than_ten_kib_is_rejected_before_extraction(self, mock_extract, tmp_path):
        source = tmp_path / "small.txt"
        source.write_bytes(b"a" * ((10 * 1024) - 1))

        with pytest.raises(FileTooSmallError):
            prepare_document_text(source)

        mock_extract.assert_not_called()

    @patch("summarization.extract_document", return_value=ExtractedDocument("Lesson text."))
    def test_file_exactly_ten_kib_continues_to_extraction(self, mock_extract, tmp_path):
        source = tmp_path / "exact.txt"
        source.write_bytes(b"a" * (10 * 1024))

        result = prepare_document_text(source)

        assert result["normalized_text"] == "Lesson text."
        mock_extract.assert_called_once_with(source)

    @patch("summarization.extract_document", return_value=ExtractedDocument("Lesson text."))
    def test_file_larger_than_ten_kib_continues_to_extraction(self, mock_extract, tmp_path):
        source = tmp_path / "large.txt"
        source.write_bytes(b"a" * ((10 * 1024) + 1))

        result = prepare_document_text(source)

        assert result["normalized_text"] == "Lesson text."
        mock_extract.assert_called_once_with(source)

    @patch("summarization.extract_document", return_value=ExtractedDocument("  Lesson   text.  "))
    def test_prepare_document_text_does_not_call_gemini(self, mock_extract):
        result = prepare_document_text("unused.txt")
        assert result["original_text"] == "  Lesson   text.  "
        assert result["normalized_text"] == "Lesson text."
        mock_extract.assert_called_once_with("unused.txt")

    @patch("summarization.summarize_chunk", return_value="Prepared summary")
    def test_prepared_text_is_not_normalized_again(self, mock_chunk):
        result = summarize_prepared_text("Original text", "Already  prepared")
        assert result["normalized_text"] == "Already  prepared"
        assert mock_chunk.call_args.args == ("Already  prepared", ())

    @patch("summarization.get_gemini_client")
    def test_summarize_chunk_calls_gemini(self, mock_get_client):
        mock_response = MagicMock()
        mock_response.text = "A short educational summary."
        mock_client = MagicMock()
        mock_client.models.generate_content.return_value = mock_response
        mock_get_client.return_value = mock_client

        result = summarize_chunk("Some educational text about mitochondria.")
        assert result == "A short educational summary."
        mock_client.models.generate_content.assert_called_once()
        config = mock_client.models.generate_content.call_args.kwargs["config"]
        assert config.automatic_function_calling.disable is True

    @patch("summarization.get_gemini_client")
    def test_generate_final_summary_combines(self, mock_get_client):
        mock_response = MagicMock()
        mock_response.text = "Final coherent summary."
        mock_client = MagicMock()
        mock_client.models.generate_content.return_value = mock_response
        mock_get_client.return_value = mock_client

        result = generate_final_summary(["Part one summary.", "Part two summary."])
        assert result == "Final coherent summary."
        mock_client.models.generate_content.assert_called_once()

    @patch("summarization.get_gemini_client")
    def test_generate_final_summary_single_skips_api(self, mock_get_client):
        # One intermediate summary should be returned as-is (no extra API call)
        result = generate_final_summary(["Only one chunk summary."])
        assert result == "Only one chunk summary."
        mock_get_client.assert_not_called()

    @patch("summarization.get_gemini_client")
    def test_timeout_is_classified_as_unavailable(self, mock_get_client):
        mock_client = MagicMock()
        mock_client.models.generate_content.side_effect = httpx.ReadTimeout("timed out")
        mock_get_client.return_value = mock_client

        with pytest.raises(GeminiUnavailableError):
            summarize_chunk("Educational content.", timeout_ms=1000)

        mock_get_client.assert_called_once_with(timeout_ms=1000)

    @patch("summarization.get_gemini_client")
    def test_quota_error_is_classified_as_unavailable(self, mock_get_client):
        mock_client = MagicMock()
        mock_client.models.generate_content.side_effect = genai_errors.ClientError(
            429,
            {"error": {"code": 429, "status": "RESOURCE_EXHAUSTED", "message": "quota"}},
        )
        mock_get_client.return_value = mock_client

        with pytest.raises(GeminiUnavailableError):
            summarize_chunk("Educational content.")

    @patch("summarization.get_gemini_client")
    def test_empty_gemini_response_is_rejected(self, mock_get_client):
        mock_client = MagicMock()
        mock_client.models.generate_content.return_value = SimpleNamespace(text="   ")
        mock_get_client.return_value = mock_client

        with pytest.raises(GeminiResponseError):
            summarize_chunk("Educational content.")

    @patch("summarization.summarize_chunk", return_value="Chunk summary")
    @patch("summarization.generate_final_summary", return_value="Final summary")
    def test_summarize_document_returns_both_texts(self, mock_final, mock_chunk):
        with patch("summarization.os.path.getsize", return_value=10 * 1024):
            result = summarize_document(SAMPLE_TXT)

        assert "original_text" in result
        assert "normalized_text" in result
        assert "summary" in result
        assert "num_chunks" in result
        assert "Photosynthesis" in result["original_text"]
        assert result["summary"] == "Final summary"
        assert result["num_chunks"] >= 1
        # Original text must still be present (not overwritten by summary)
        assert result["original_text"] != result["summary"]

    @patch("summarization.extract_document", return_value=ExtractedDocument("  مرحبا   بالعالم!\n\n\nNext  line. "))
    @patch("summarization.summarize_chunk", return_value="Summary")
    def test_normalized_text_is_sent_to_model(self, mock_chunk, mock_extract):
        result = summarize_document("unused.txt")
        assert result["original_text"] == "  مرحبا   بالعالم!\n\n\nNext  line. "
        assert result["normalized_text"] == "مرحبا بالعالم!\n\nNext line."
        assert mock_chunk.call_args.args == (result["normalized_text"], ())
        assert 0 < mock_chunk.call_args.kwargs["timeout_ms"] <= 60_000

    @patch("summarization.extract_document", return_value=ExtractedDocument(" \t\n\n "))
    @patch("summarization.summarize_chunk")
    def test_empty_normalized_text_is_rejected(self, mock_chunk, mock_extract):
        with pytest.raises(ValueError, match="No readable text"):
            summarize_document("unused.txt")
        mock_chunk.assert_not_called()

    @patch(
        "summarization.extract_document",
        return_value=ExtractedDocument("Contents\nChapter 1 .... 4\n\nChapter 1: Cells\nCells have nuclei.\nPage 4"),
    )
    @patch("summarization.summarize_chunk", return_value="Summary")
    def test_structure_is_context_but_not_summary_content(self, mock_chunk, mock_extract):
        result = summarize_document("unused.txt")
        assert result["original_text"].startswith("Contents")
        assert result["normalized_text"] == "Chapter 1: Cells\nCells have nuclei."
        assert result["structure"]["page_numbers"] == ["4"]
        assert result["structure"]["headings"] == ["Chapter 1: Cells"]
        assert result["structure"]["table_of_contents"] == [{"title": "Chapter 1", "page_number": "4"}]
        assert mock_chunk.call_args.args == (result["normalized_text"], ("Chapter 1: Cells",))
        assert 0 < mock_chunk.call_args.kwargs["timeout_ms"] <= 60_000


class TestApiKeyHandling:
    def test_missing_api_key_raises_clear_error(self, monkeypatch):
        monkeypatch.delenv("GEMINI_API_KEY", raising=False)
        with pytest.raises(RuntimeError, match="not configured"):
            summarization.get_gemini_client()


class TestFlaskApi:
    def setup_method(self):
        self.client = app.test_client()
        os.environ["AI_SERVICE_INTERNAL_KEY"] = "test-only-shared-key"

    @patch("summarization_api.summarize_document")
    def test_summarize_endpoint_success(self, mock_summarize):
        mock_summarize.return_value = {
            "original_text": "Original educational content.",
            "normalized_text": "Original educational content.",
            "summary": "Short summary.",
            "num_chunks": 1,
        }

        with open(SAMPLE_TXT, "rb") as f:
            response = self.client.post(
                "/api/summarize",
                data={"file": (f, "sample.txt")},
                content_type="multipart/form-data",
            )

        assert response.status_code == 200
        data = response.get_json()
        assert data["success"] is True
        assert data["filename"] == "sample.txt"
        assert data["original_text"] == "Original educational content."
        assert data["normalized_text"] == "Original educational content."
        assert data["summary"] == "Short summary."
        assert data["num_chunks"] == 1

    def test_summarize_endpoint_rejects_file_smaller_than_ten_kib(self):
        response = self.client.post(
            "/api/summarize",
            data={"file": (BytesIO(b"a" * ((10 * 1024) - 1)), "small.txt")},
            content_type="multipart/form-data",
        )

        assert response.status_code == 422
        assert response.get_json() == {
            "success": False,
            "error": "The uploaded file must be at least 10 KB to summarize.",
            "errorAr": "يجب ألا يقل حجم الملف المرفوع عن 10 كيلوبايت لإجراء التلخيص.",
        }

    @patch("summarization_api.prepare_document_text")
    def test_extract_text_endpoint_does_not_summarize(self, mock_prepare):
        mock_prepare.return_value = {
            "original_text": "Original educational content.",
            "normalized_text": "Original educational content.",
            "structure": {},
        }
        with open(SAMPLE_TXT, "rb") as file:
            response = self.client.post(
                "/api/extract-text",
                data={"file": (file, "sample.txt")},
                content_type="multipart/form-data",
                headers={"X-Internal-Api-Key": "test-only-shared-key"},
            )
        assert response.status_code == 200
        assert response.get_json() == {
            "success": True,
            "original_text": "Original educational content.",
            "normalized_text": "Original educational content.",
            "structure": {},
        }
        mock_prepare.assert_called_once()

    def test_extract_text_endpoint_requires_internal_key(self):
        response = self.client.post(
            "/api/extract-text",
            data={"file": (BytesIO(b"Study text"), "sample.txt")},
            content_type="multipart/form-data",
            headers={"X-Internal-Api-Key": "wrong-key"},
        )
        assert response.status_code == 401

    @patch("summarization_api.summarize_prepared_text")
    def test_summarize_text_endpoint_reuses_prepared_text(self, mock_summarize):
        mock_summarize.return_value = {
            "original_text": "Original educational content.",
            "normalized_text": "Normalized educational content.",
            "summary": "Short summary.",
            "num_chunks": 1,
            "structure": {},
        }
        response = self.client.post(
            "/api/summarize-text",
            headers={"X-Internal-Api-Key": "test-only-shared-key"},
            json={
                "original_text": "Original educational content.",
                "normalized_text": "Normalized educational content.",
            },
        )
        assert response.status_code == 200
        assert response.get_json()["summary"] == "Short summary."
        mock_summarize.assert_called_once_with(
            "Original educational content.",
            "Normalized educational content.",
            None,
        )

    def test_summarize_text_endpoint_requires_internal_key(self):
        response = self.client.post(
            "/api/summarize-text",
            headers={"X-Internal-Api-Key": "wrong-key"},
            json={
                "original_text": "Original educational content.",
                "normalized_text": "Normalized educational content.",
            },
        )
        assert response.status_code == 401

    def test_summarize_endpoint_no_file(self):
        response = self.client.post("/api/summarize")
        assert response.status_code == 400
        data = response.get_json()
        assert data["success"] is False

    @patch("summarization_api.summarize_document", side_effect=DocConverterUnavailableError("missing"))
    def test_doc_without_converter_returns_clear_json(self, mock_summarize):
        response = self.client.post(
            "/api/summarize",
            data={"file": (BytesIO(b"legacy Word data"), "lesson.doc")},
            content_type="multipart/form-data",
        )
        assert response.status_code == 503
        assert response.get_json() == {
            "success": False,
            "code": "LIBREOFFICE_REQUIRED",
            "error": "LibreOffice must be installed to read .doc files.",
        }

    @patch("summarization_api.summarize_document", side_effect=EmptyDocumentError("empty"))
    def test_summarize_endpoint_empty_text(self, mock_summarize):
        with open(SAMPLE_TXT, "rb") as f:
            response = self.client.post(
                "/api/summarize", data={"file": (f, "sample.txt")}, content_type="multipart/form-data"
            )
        assert response.status_code == 422
        assert response.get_json()["error"] == "No readable text was found in the document."

    @patch("summarization_api.summarize_document", side_effect=GeminiConfigurationError("missing key"))
    def test_summarize_endpoint_missing_key_returns_503(self, mock_summarize):
        with open(SAMPLE_TXT, "rb") as file:
            response = self.client.post(
                "/api/summarize", data={"file": (file, "sample.txt")}, content_type="multipart/form-data"
            )
        assert response.status_code == 503
        assert "missing key" not in response.get_data(as_text=True)

    @patch("summarization_api.summarize_document", side_effect=GeminiUnavailableError("timeout"))
    def test_summarize_endpoint_timeout_returns_503(self, mock_summarize):
        with open(SAMPLE_TXT, "rb") as file:
            response = self.client.post(
                "/api/summarize", data={"file": (file, "sample.txt")}, content_type="multipart/form-data"
            )
        assert response.status_code == 503
        assert "timeout" not in response.get_data(as_text=True)

    @patch("summarization_api.summarize_document", side_effect=GeminiResponseError("empty response"))
    def test_summarize_endpoint_empty_model_response_returns_502(self, mock_summarize):
        with open(SAMPLE_TXT, "rb") as file:
            response = self.client.post(
                "/api/summarize", data={"file": (file, "sample.txt")}, content_type="multipart/form-data"
            )
        assert response.status_code == 502
        assert "empty response" not in response.get_data(as_text=True)

    @patch("summarization_api.summarize_document")
    def test_summarize_endpoint_malformed_result_returns_502(self, mock_summarize):
        mock_summarize.return_value = {
            "original_text": "Original text",
            "normalized_text": "Original text",
            "summary": "",
            "num_chunks": 1,
        }
        with open(SAMPLE_TXT, "rb") as file:
            response = self.client.post(
                "/api/summarize", data={"file": (file, "sample.txt")}, content_type="multipart/form-data"
            )
        assert response.status_code == 502

    @patch("summarization_api.summarize_document", side_effect=KeyError("private internal detail"))
    def test_summarize_endpoint_unexpected_error_logs_and_returns_safe_500(self, mock_summarize):
        with patch.object(summarization_api.logger, "exception") as log_exception:
            with open(SAMPLE_TXT, "rb") as file:
                response = self.client.post(
                    "/api/summarize", data={"file": (file, "sample.txt")}, content_type="multipart/form-data"
                )
        assert response.status_code == 500
        assert "private internal detail" not in response.get_data(as_text=True)
        log_exception.assert_called_once_with("Unexpected summarization failure: %s", "KeyError")

    def test_grading_internal_key_rejects_unauthorized_request(self, monkeypatch):
        monkeypatch.setenv("AI_SERVICE_INTERNAL_KEY", "test-only-shared-key")
        response = self.client.post(
            "/api/grade-short-answer",
            json={"question": "Q", "student_answer": "A", "difficulty": "EASY", "source_text": "Source"},
            headers={"X-Internal-Api-Key": "wrong-key"},
        )
        assert response.status_code == 401

    @patch("summarization_api.grade_short_answer", return_value={
        "isCorrect": True, "scoreRatio": 1.0, "feedback": "Correct", "needsReview": False,
    })
    def test_grading_internal_key_accepts_matching_request(self, mock_grade, monkeypatch):
        monkeypatch.setenv("AI_SERVICE_INTERNAL_KEY", "test-only-shared-key")
        response = self.client.post(
            "/api/grade-short-answer",
            json={"question": "Q", "student_answer": "A", "difficulty": "EASY", "source_text": "Source"},
            headers={"X-Internal-Api-Key": "test-only-shared-key"},
        )
        assert response.status_code == 200
        mock_grade.assert_called_once()


# ---------------------------------------------------------------------------
# Integration tests (REAL Gemini API + real files)
# ---------------------------------------------------------------------------


@skip_without_key
class TestRealGeminiIntegration:
    """
    These tests call the real Gemini model.
    They print the generated summary so you can visually check quality.
    """

    def _run_and_print(self, path, label):
        print(f"\n{'=' * 60}")
        print(f"INTEGRATION TEST: {label}")
        print(f"File: {path}")
        print("=" * 60)

        result = summarize_document(path)

        print(f"\nnum_chunks: {result['num_chunks']}")
        print(f"\n--- ORIGINAL TEXT (first 400 chars) ---")
        print(result["original_text"][:400])
        print(f"\n--- GENERATED SUMMARY ---")
        print(result["summary"])
        print("=" * 60)

        assert result["original_text"]
        assert result["summary"]
        assert result["num_chunks"] >= 1
        # Original must still be intact for the teammate's question-generation task
        assert "Photosynthesis" in result["original_text"] or "photosynthesis" in result[
            "original_text"
        ].lower()
        return result

    def test_real_summary_txt(self):
        self._run_and_print(SAMPLE_TXT, "TXT file")

    def test_real_summary_pdf(self):
        self._run_and_print(SAMPLE_PDF, "PDF file")

    def test_real_summary_docx(self):
        self._run_and_print(SAMPLE_DOCX, "DOCX file")
